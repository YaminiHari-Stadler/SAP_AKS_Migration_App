"""Turn migration entries into artefacts a person can read.

The engine speaks in COPIED / CONVERTED / MANUAL_REVIEW / SKIPPED. Engineers
reading the result want OK / REVIEW / WARNING / ERROR. That translation lives
here, as a presentation layer only: no value written into the migrated workbook
is affected by anything in this module, and `write_csv_report` still produces
the exact V1 pilot audit format.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import IDENTITY_KEYS, REPORT_COLUMNS, MigrationEntry, display

STATUS_OK = "OK"
STATUS_REVIEW = "REVIEW"
STATUS_WARNING = "WARNING"
STATUS_ERROR = "ERROR"
STATUS_SKIPPED = "SKIPPED"

REVIEW_STATUSES = (STATUS_REVIEW, STATUS_WARNING, STATUS_ERROR)

TABLE_COLUMNS = [
    "Status",
    "Source Row",
    "Position",
    "Equipment",
    "Parameter",
    "SAP Field",
    "Old Value",
    "New Value",
    "Reason",
    "Source Cell",
    "Target Cell",
    "Written",
]

STATUS_FILLS = {
    STATUS_OK: "C6EFCE",
    STATUS_REVIEW: "FFC7CE",
    STATUS_WARNING: "FFEB9C",
    STATUS_ERROR: "F8CBAD",
    STATUS_SKIPPED: "EDEDED",
}


def friendly_status(entry: MigrationEntry) -> str:
    """Project an engine status onto the vocabulary used in the user interface."""
    if entry.status == "SKIPPED":
        return STATUS_SKIPPED
    if entry.status == "MANUAL_REVIEW":
        return STATUS_REVIEW
    note = (entry.note or "").casefold()
    if "close header match" in note or "status-code label matched at" in note:
        return STATUS_WARNING
    return STATUS_OK


def reason_for(entry: MigrationEntry) -> str:
    """A plain-language explanation of why an entry ended up with its status."""
    note = entry.note or ""
    low = note.casefold()

    if entry.status == "SKIPPED":
        return "The legacy cell is empty, so nothing is written."
    if "target left blank because multiple non-equivalent values" in low:
        return (
            "Two legacy columns feed this SAP field with different values. "
            "Left blank — the engineering team decides which one applies."
        )
    if "conflicting source fields map to the same sap target" in low:
        return "This value conflicts with another legacy column writing to the same SAP field. Not written."
    if "status code not resolved safely" in low:
        return "No numeric status code is defined for this status text."
    if "not listed in the template's allowed sap values" in low:
        if entry.has_mapping_table:
            return "No migration mapping is defined for this old value."
        return "This value is not in the AKS V2 list of allowed values for this parameter."
    if "duplicate target value not written twice" in low:
        return "Another legacy column already wrote the same value to this cell."
    if entry.status == "CONVERTED":
        head = note.split(";")[0].strip()
        if "→" in head:
            return f"Converted by a mapping rule: {head}."
        if head:
            return f"Converted: {head}."
        return "Converted by a mapping rule."
    if "close header match" in low:
        return (
            "Copied unchanged, but the legacy column heading only matched approximately. "
            "Confirm that the right column was read."
        )
    if "status-code label matched at" in low:
        return f"Status code derived from an approximate label match ({note.split(';')[0].strip()})."
    return "Copied unchanged."


def row_identities(entries: Iterable[MigrationEntry]) -> dict[int, dict[str, str]]:
    """Map each legacy row number to the values that identify it for a human."""
    identities: dict[int, dict[str, str]] = {}
    for entry in entries:
        if entry.target_key not in IDENTITY_KEYS:
            continue
        bucket = identities.setdefault(entry.source_row, {})
        value = display(entry.old_value).strip()
        if value and not bucket.get(entry.target_key):
            bucket[entry.target_key] = value
    resolved: dict[int, dict[str, str]] = {}
    for source_row, bucket in identities.items():
        resolved[source_row] = {
            "position": bucket.get("CUST_SPEC_DESI") or bucket.get("PROC_POSI_NUMB") or "",
            "equipment": bucket.get("TITLE") or "",
        }
    return resolved


def to_table_rows(
    entries: Sequence[MigrationEntry],
    identities: dict[int, dict[str, str]] | None = None,
    statuses: Sequence[str] | None = None,
) -> list[dict[str, Any]]:
    """Build display rows, optionally filtered to certain friendly statuses."""
    identities = identities if identities is not None else row_identities(entries)
    wanted = set(statuses) if statuses else None
    rows: list[dict[str, Any]] = []
    for entry in entries:
        status = friendly_status(entry)
        if wanted is not None and status not in wanted:
            continue
        identity = identities.get(entry.source_row, {})
        rows.append(
            {
                "Status": status,
                "Source Row": entry.source_row,
                "Position": identity.get("position", ""),
                "Equipment": identity.get("equipment", ""),
                "Parameter": entry.target_label or entry.target_key,
                "SAP Field": entry.target_key,
                "Old Value": display(entry.old_value),
                "New Value": display(entry.new_value),
                "Reason": reason_for(entry),
                "Source Cell": entry.source_cell,
                "Target Cell": entry.target_cell,
                "Written": "yes" if (entry.apply and entry.status != "SKIPPED") else "no",
            }
        )
    return rows


def count_statuses(entries: Iterable[MigrationEntry]) -> dict[str, int]:
    counts = {
        STATUS_OK: 0,
        STATUS_REVIEW: 0,
        STATUS_WARNING: 0,
        STATUS_ERROR: 0,
        STATUS_SKIPPED: 0,
    }
    for entry in entries:
        counts[friendly_status(entry)] += 1
    return counts


def review_breakdown(entries: Sequence[MigrationEntry]) -> list[dict[str, Any]]:
    """Group review and warning items by parameter so a reviewer sees the themes."""
    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}
    for entry in entries:
        status = friendly_status(entry)
        if status not in REVIEW_STATUSES:
            continue
        reason = reason_for(entry)
        key = (entry.target_key, entry.target_label or entry.target_key, reason)
        bucket = grouped.setdefault(
            key,
            {
                "SAP Field": entry.target_key,
                "Parameter": entry.target_label or entry.target_key,
                "Reason": reason,
                "Status": status,
                "Items": 0,
                "_examples": [],
            },
        )
        bucket["Items"] += 1
        old = display(entry.old_value).strip()
        if old and old not in bucket["_examples"] and len(bucket["_examples"]) < 5:
            bucket["_examples"].append(old)
    rows = []
    for bucket in grouped.values():
        examples = bucket.pop("_examples")
        bucket["Example old values"] = " | ".join(examples)
        rows.append(bucket)
    rows.sort(key=lambda row: row["Items"], reverse=True)
    return rows


# ---------------------------------------------------------------------------
# File artefacts
# ---------------------------------------------------------------------------


def write_csv_report(path: Path, entries: Iterable[MigrationEntry]) -> Path:
    """Write the V1 pilot audit CSV, unchanged, for comparison with earlier runs."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(REPORT_COLUMNS)
        for entry in entries:
            writer.writerow([display(getattr(entry, column)) for column in REPORT_COLUMNS])
    return path


def _style_header(sheet, columns: Sequence[str], row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row, index)
        cell.value = name
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")


def write_audit_xlsx(
    path: Path,
    entries: Sequence[MigrationEntry],
    identities: dict[int, dict[str, str]] | None = None,
) -> Path:
    """Full cell-level before/after audit, one row per mapped value."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = to_table_rows(entries, identities)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Audit"
    _style_header(sheet, TABLE_COLUMNS)
    fills = {status: PatternFill("solid", fgColor=color) for status, color in STATUS_FILLS.items()}
    for row in rows:
        sheet.append([row[column] for column in TABLE_COLUMNS])
        fill = fills.get(row["Status"])
        if fill is not None:
            sheet.cell(sheet.max_row, 1).fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(TABLE_COLUMNS))}{sheet.max_row}"
    widths = [10, 11, 12, 34, 26, 18, 28, 28, 60, 12, 12, 9]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(path)
    return path


def write_report_xlsx(path: Path, context: dict[str, Any], entries: Sequence[MigrationEntry]) -> Path:
    """Human-readable migration report: what ran, what came out, what needs a decision."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    summary = context.get("summary", {})
    counts = context.get("counts", {})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Migration Report"

    title = sheet.cell(1, 1)
    title.value = "AKS Migration Report"
    title.font = Font(bold=True, size=16, color="1F4E79")

    facts: list[tuple[str, Any]] = [
        ("Migration timestamp", context.get("timestamp", datetime.now().isoformat(timespec="seconds"))),
        ("Run ID", context.get("run_id", "")),
        ("Result", context.get("result", "")),
        ("", ""),
        ("Source file", context.get("source_name", "")),
        ("Source worksheet", context.get("source_sheet", "")),
        ("Mapping file (version)", context.get("mapping_name", "")),
        ("AKS V2 template (version)", context.get("template_name", "")),
        ("Migrated workbook", context.get("output_name", "")),
        ("", ""),
        ("Rows processed", summary.get("source_rows", 0)),
        ("Empty legacy rows skipped", summary.get("placeholder_rows", 0)),
        ("Active mapping rules", summary.get("active_rules", 0)),
        ("Values migrated", summary.get("applied_writes", 0)),
        ("Mapped values examined", summary.get("entries", 0)),
        ("", ""),
        ("Successful mappings (OK)", counts.get(STATUS_OK, 0)),
        ("Warnings", counts.get(STATUS_WARNING, 0)),
        ("Items requiring review", counts.get(STATUS_REVIEW, 0)),
        ("Errors", counts.get(STATUS_ERROR, 0)),
        ("Empty source cells skipped", counts.get(STATUS_SKIPPED, 0)),
        ("", ""),
        ("Legacy rows read", f"{summary.get('first_source_row', '')} – {summary.get('last_source_row', '')}"),
        ("Target rows written", f"{summary.get('first_target_row', '')} – {summary.get('last_target_row', '')}"),
        ("Output verified after writing", context.get("verification", "")),
        ("Log file", context.get("log_name", "")),
    ]
    row_number = 3
    for label, value in facts:
        if label:
            key_cell = sheet.cell(row_number, 1)
            key_cell.value = label
            key_cell.font = Font(bold=True)
            sheet.cell(row_number, 2).value = value
        row_number += 1

    row_number += 1
    heading = sheet.cell(row_number, 1)
    heading.value = "Warnings and notes"
    heading.font = Font(bold=True, size=12)
    row_number += 1
    notes = list(context.get("warnings", []))
    if not notes:
        notes = ["No warnings were raised."]
    for note in notes:
        sheet.cell(row_number, 1).value = "•"
        cell = sheet.cell(row_number, 2)
        cell.value = note
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_number += 1

    errors = list(context.get("errors", []))
    if errors:
        row_number += 1
        heading = sheet.cell(row_number, 1)
        heading.value = "Errors"
        heading.font = Font(bold=True, size=12, color="C00000")
        row_number += 1
        for message in errors:
            sheet.cell(row_number, 1).value = "•"
            cell = sheet.cell(row_number, 2)
            cell.value = message
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            row_number += 1

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 96

    breakdown = review_breakdown(entries)
    unresolved = workbook.create_sheet("Unresolved mappings")
    columns = ["Status", "Parameter", "SAP Field", "Reason", "Items", "Example old values"]
    _style_header(unresolved, columns)
    if breakdown:
        for row in breakdown:
            unresolved.append([row.get(column, "") for column in columns])
    else:
        unresolved.append(["OK", "", "", "Every value was covered by the migration mapping.", 0, ""])
    unresolved.freeze_panes = "A2"
    unresolved.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{unresolved.max_row}"
    for index, width in enumerate([10, 30, 18, 62, 9, 52], start=1):
        unresolved.column_dimensions[get_column_letter(index)].width = width

    workbook.save(path)
    return path
