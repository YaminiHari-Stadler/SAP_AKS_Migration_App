"""Pre-flight and post-flight safety checks.

Nothing here changes a file. These functions answer three questions:

1. can this machine run a migration at all (Excel + COM automation)?
2. are the three input workbooks the ones this tool expects?
3. after Excel has written the output, does the output actually contain the
   rows we asked for?
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from .excel_reader import load_sheet, sheet_names
from .model import (
    MAPPING_SHEET,
    SOURCE_DATA_START_ROW,
    SOURCE_HEADER_ROWS,
    SOURCE_SHEET,
    TARGET_DATA_START_ROW,
    TARGET_KEY_ROW,
    TARGET_SHEET,
    VALUES_SHEET,
    MigrationError,
    display,
    is_blank,
)

# Columns that make a target row count as populated (project number, position,
# designation, type). Kept in sync with the template-demo-data detection.
POPULATED_TARGET_COLUMNS = (3, 4, 12, 13)


@dataclass
class Check:
    name: str
    ok: bool
    detail: str


@dataclass
class EnvironmentReport:
    ok: bool
    checks: list[Check] = field(default_factory=list)

    @property
    def failures(self) -> list[Check]:
        return [check for check in self.checks if not check.ok]


@dataclass
class SourceInspection:
    path: Path
    ok: bool
    sheet: str | None = None
    available_sheets: list[str] = field(default_factory=list)
    data_rows: int = 0
    placeholder_rows: int = 0
    first_row: int | None = None
    last_row: int | None = None
    header_rows_found: list[int] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class OutputVerification:
    ok: bool
    output_path: Path
    size_bytes: int = 0
    expected_first_row: int = 0
    expected_last_row: int = 0
    actual_last_row: int | None = None
    populated_rows: int = 0
    report_sheet_present: bool = False
    problems: list[str] = field(default_factory=list)


def check_environment() -> EnvironmentReport:
    """Verify that this machine can run the Excel write phase."""
    checks: list[Check] = []

    try:
        import pyxlsb  # noqa: F401
        checks.append(Check("Legacy .xlsb reader (pyxlsb)", True, "available"))
    except Exception as exc:
        checks.append(Check("Legacy .xlsb reader (pyxlsb)", False, f"not installed: {exc}"))

    try:
        import openpyxl as _openpyxl
        checks.append(Check("Excel .xlsx reader (openpyxl)", True, f"version {_openpyxl.__version__}"))
    except Exception as exc:
        checks.append(Check("Excel .xlsx reader (openpyxl)", False, f"not installed: {exc}"))

    try:
        import pythoncom  # noqa: F401
        import win32com.client  # noqa: F401
        checks.append(Check("COM automation (pywin32)", True, "available"))
    except Exception as exc:
        checks.append(Check("COM automation (pywin32)", False, f"not installed: {exc}"))
        checks.append(Check("Microsoft Excel", False, "cannot be checked without pywin32"))
        return EnvironmentReport(ok=False, checks=checks)

    excel = None
    try:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        excel = win32com.client.DispatchEx("Excel.Application")
        version = str(excel.Version)
        checks.append(Check("Microsoft Excel", True, f"desktop Excel {version} responded to automation"))
    except Exception as exc:
        checks.append(
            Check(
                "Microsoft Excel",
                False,
                "Excel did not start through automation. Install Microsoft Excel for Windows "
                f"(desktop, not the web version) and try again. Technical detail: {exc}",
            )
        )
    finally:
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
        try:
            import pythoncom

            pythoncom.CoUninitialize()
        except Exception:
            pass

    return EnvironmentReport(ok=all(check.ok for check in checks), checks=checks)


def inspect_source(path: Path) -> SourceInspection:
    """Confirm a file really is a legacy AKS workbook and count its equipment rows."""
    path = Path(path)
    result = SourceInspection(path=path, ok=False)

    if not path.exists() or not path.is_file():
        result.problems.append(f"File not found: {path}")
        return result
    if path.suffix.lower() != ".xlsb":
        result.problems.append(
            f"{path.name} is not an Excel binary workbook. Legacy AKS lists are .xlsb files."
        )
        return result

    try:
        result.available_sheets = sheet_names(path)
    except MigrationError as exc:
        result.problems.append(str(exc))
        return result

    if SOURCE_SHEET not in result.available_sheets:
        result.problems.append(
            f"This workbook has no '{SOURCE_SHEET}' worksheet, so it is not a legacy AKS list. "
            f"Worksheets found: {', '.join(result.available_sheets) or 'none'}."
        )
        return result
    result.sheet = SOURCE_SHEET

    try:
        rows = load_sheet(path, SOURCE_SHEET)
    except MigrationError as exc:
        result.problems.append(str(exc))
        return result

    result.header_rows_found = [row for row in SOURCE_HEADER_ROWS if rows.get(row)]
    if not result.header_rows_found:
        result.problems.append(
            "No column headings were found on rows "
            f"{', '.join(str(row) for row in SOURCE_HEADER_ROWS)} of '{SOURCE_SHEET}'. "
            "This does not look like a legacy AKS list."
        )
        return result

    # Local import avoids a circular import at module load time.
    from .engine import source_data_rows

    try:
        data_rows, placeholders = source_data_rows(rows)
    except MigrationError as exc:
        result.problems.append(str(exc))
        return result

    result.data_rows = len(data_rows)
    result.placeholder_rows = placeholders
    result.first_row = data_rows[0]
    result.last_row = data_rows[-1]
    if placeholders:
        result.notes.append(
            f"{placeholders} rows below row {SOURCE_DATA_START_ROW} hold only pre-filled defaults "
            "and are not treated as equipment."
        )
    result.ok = True
    return result


def validate_template(path: Path) -> list[str]:
    """Return a list of problems with the AKS V2 template, empty when it is usable."""
    path = Path(path)
    problems: list[str] = []
    if not path.exists():
        return [f"AKS V2 template not found: {path}"]
    if path.suffix.lower() != ".xlsb":
        return [f"The AKS V2 template must be an .xlsb file, got {path.name}"]
    try:
        names = sheet_names(path)
    except MigrationError as exc:
        return [str(exc)]
    for required in (TARGET_SHEET, VALUES_SHEET):
        if required not in names:
            problems.append(f"The template has no '{required}' worksheet.")
    if problems:
        return problems
    try:
        rows = load_sheet(path, TARGET_SHEET)
    except MigrationError as exc:
        return [str(exc)]
    keys = [value for value in rows.get(TARGET_KEY_ROW, {}).values() if not is_blank(value)]
    if len(keys) < 10:
        problems.append(
            f"Row {TARGET_KEY_ROW} of the template's '{TARGET_SHEET}' sheet should hold the SAP "
            f"field keys, but only {len(keys)} were found. The template version may be wrong."
        )
    return problems


def validate_mapping(path: Path) -> list[str]:
    """Return a list of problems with the migration mapping workbook."""
    path = Path(path)
    if not path.exists():
        return [f"Migration mapping file not found: {path}"]
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return [f"The migration mapping must be an .xlsx file, got {path.name}"]
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return [f"{path.name} could not be opened as an Excel workbook: {exc}"]
    try:
        if MAPPING_SHEET not in workbook.sheetnames:
            return [
                f"The mapping workbook has no '{MAPPING_SHEET}' worksheet. "
                f"Worksheets found: {', '.join(workbook.sheetnames)}."
            ]
    finally:
        workbook.close()
    return []


def verify_output(
    output_path: Path, summary: dict[str, Any]
) -> OutputVerification:
    """Reopen the generated workbook and confirm the migrated rows are really there."""
    output_path = Path(output_path)
    expected_first = int(summary.get("first_target_row", TARGET_DATA_START_ROW))
    expected_last = int(summary.get("last_target_row", TARGET_DATA_START_ROW))
    verification = OutputVerification(
        ok=False,
        output_path=output_path,
        expected_first_row=expected_first,
        expected_last_row=expected_last,
    )

    if not output_path.exists():
        verification.problems.append("The migrated workbook was not created.")
        return verification
    verification.size_bytes = output_path.stat().st_size
    if verification.size_bytes == 0:
        verification.problems.append("The migrated workbook is empty (0 bytes).")
        return verification

    try:
        names = sheet_names(output_path)
        verification.report_sheet_present = "Migration_Report" in names
        # The file was just written, so never serve it from the read cache.
        rows = load_sheet(output_path, TARGET_SHEET, use_cache=False)
    except MigrationError as exc:
        verification.problems.append(f"The migrated workbook could not be reopened: {exc}")
        return verification

    populated = [
        row_number
        for row_number, values in rows.items()
        if row_number >= expected_first
        and any(not is_blank(values.get(column)) for column in POPULATED_TARGET_COLUMNS)
    ]
    verification.populated_rows = len(populated)
    verification.actual_last_row = max(populated) if populated else None

    if not populated:
        verification.problems.append(
            f"No migrated data rows were found from row {expected_first} downwards."
        )
    elif verification.actual_last_row != expected_last:
        verification.problems.append(
            f"The last populated row is {verification.actual_last_row}, "
            f"but row {expected_last} was expected."
        )
    if not verification.report_sheet_present:
        verification.problems.append("The 'Migration_Report' worksheet is missing from the output.")

    verification.ok = not verification.problems
    return verification
